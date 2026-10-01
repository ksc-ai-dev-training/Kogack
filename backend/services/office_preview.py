# 参照ドキュメントのOffice形式（Word・PowerPoint・Excel）のプレビューと索引用テキスト抽出。
# Word（.docx）のプレビュー（ユーザー要望、2026-10-01）として作り、同日にPowerPoint（.pptx）・
# Excel（.xlsx）の対応（ユーザー要望「ExcelとPowerPointも読めるようにすべきか」→実施）を追加した
# （旧docx_preview.py）。
# ブラウザはOffice形式をそのまま表示できないため、サーバー側で本文を見出し・段落・箇条書き・表の
# ブロック列（JSON）へ変換して返す。HTMLへの変換（mammoth等）ではなく構造化データにしているのは、
# フロント側がReactのテキストとして描画するだけで済み、文書内に仕込まれたリンク・スクリプト等を
# ブラウザに解釈させる余地を作らないため。文字装飾・画像・グラフ・段組み等は再現しない（内容の
# 確認用途に絞る）。
# PowerPoint・Excelの索引化（doc_indexer.py）も同じブロック列からテキストを作る（blocks_to_text）。
# プレビューと索引で別々に読み取り処理を持つと、AIが読んでいる内容と画面で確認できる内容が
# 食い違いうるため一本化した（Wordの索引化は従来どおり段落のみを抽出しており、ここでは変えていない）。
import io
import re
from datetime import date, datetime, time

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# 極端に大きな文書でレスポンスが膨らむのを防ぐ上限（超過分は打ち切り、truncated=Trueで知らせる）
MAX_BLOCKS = 3000
# Excelのプレビューで1シートあたりに表示する行数・列数の上限（索引化では行数の上限を設けない。
# 索引側はdoc_indexer.MAX_CHUNKSで別途打ち切られる）
XLSX_PREVIEW_MAX_ROWS = 500
XLSX_MAX_COLS = 50
OFFICE_EXTENSIONS = {".docx", ".pptx", ".xlsx"}


def _paragraph_block(p: Paragraph) -> dict | None:
    text = p.text
    style = (p.style.name if p.style is not None else "") or ""
    if style == "Title":
        return {"type": "heading", "level": 1, "text": text}
    if style.startswith("Heading"):
        level = style.removeprefix("Heading").strip()
        return {"type": "heading", "level": min(int(level), 4) if level.isdigit() else 2, "text": text}
    # 箇条書き・番号付きリストはスタイル名（List Bullet等）か段落の番号設定（numPr）で判定する
    is_list = style.startswith("List") or p._p.pPr is not None and p._p.pPr.numPr is not None
    if is_list and text.strip():
        return {"type": "list_item", "text": text}
    return {"type": "paragraph", "text": text}


def _table_block(t: Table) -> dict:
    rows = []
    for row in t.rows:
        cells: list[dict] = []
        prev = None
        for cell in row.cells:
            # 横方向の結合セルはpython-docxでは同じセルが繰り返し返されるため、連続する重複は
            # 1つにまとめてspan（結合した列数）を数える。まとめないと列がずれ、単に捨てると
            # 行ごとの列数が揃わず表が崩れる
            if prev is not None and cell._tc is prev:
                cells[-1]["span"] += 1
                continue
            prev = cell._tc
            cells.append({"text": cell.text, "span": 1})
        rows.append(cells)
    return {"type": "table", "rows": rows}


def _docx_to_blocks(data: bytes) -> dict:
    doc = DocxDocument(io.BytesIO(data))
    blocks: list[dict] = []
    truncated = False
    # doc.paragraphsとdoc.tablesは別々のリストで本文中の順序が失われるため、body要素を順に辿る
    for child in doc.element.body.iterchildren():
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        if child.tag.endswith("}p"):
            block = _paragraph_block(Paragraph(child, doc))
        elif child.tag.endswith("}tbl"):
            block = _table_block(Table(child, doc))
        else:
            continue
        if block is not None:
            blocks.append(block)
    # 末尾の空段落は表示上の意味が無いので落とす
    while blocks and blocks[-1]["type"] == "paragraph" and not blocks[-1]["text"].strip():
        blocks.pop()
    return {"blocks": blocks, "truncated": truncated}


# PowerPointが画像に自動で付ける代替テキスト（「テキスト が含まれている画像 自動的に生成された説明」等）。
# 内容をほとんど表さないため、AIへ渡すとノイズになるだけなので除外する
_AUTO_ALT_TEXT_MARKERS = ("自動的に生成された説明", "automatically generated")
_IMAGE_FILE_NAME = re.compile(r"[^/\\]+\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)", re.IGNORECASE)


def _chart_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _pptx_chart_blocks(chart) -> list[dict]:
    """グラフ（2026-10-01、ユーザー要望で追加）。PowerPointのグラフは描画用の画像ではなく元データ
    （タイトル・項目・系列ごとの値）をファイル内に持っているため、それを表として取り出す。
    横軸の項目を行、系列を列にする。散布図など項目を持たない種類は1, 2, 3…の連番で代用する。
    想定外の構造のグラフで読み取りに失敗しても、スライドの他の内容の読み取りは続ける"""
    try:
        title = chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
        series = [s for plot in chart.plots for s in plot.series]
        if not series:
            return []
        try:
            categories = [str(c) for c in chart.plots[0].categories]
        except Exception:
            categories = []
        # 散布図は項目（categories）を持たず、X値はpython-pptxの公開APIに無いためXMLのキャッシュから読む
        x_label = "項目"
        if not categories:
            xs = series[0]._element.xpath("./c:xVal//c:pt/c:v/text()")
            if xs:
                categories = [_chart_value(float(x)) if x.replace(".", "", 1).lstrip("-").isdigit() else x for x in xs]
                x_label = "X"
        n = max(len(categories), max(len(list(s.values)) for s in series))
        categories = (categories + [str(i + 1) for i in range(len(categories), n)])[:n]
        header = [x_label] + [s.name or f"系列{i + 1}" for i, s in enumerate(series)]
        values = [list(s.values) for s in series]
        rows = [header] + [
            [categories[r]] + [_chart_value(v[r]) if r < len(v) else "" for v in values] for r in range(n)
        ]
    except Exception:
        return [{"type": "paragraph", "text": "グラフ（内容を読み取れませんでした）"}]
    return [
        {"type": "paragraph", "text": f"グラフ: {title}" if title else "グラフ"},
        # has_header: 1行目が必ず見出し（項目・系列名）であることを索引化（_table_to_text）へ伝える
        {"type": "table", "rows": [[{"text": c, "span": 1} for c in r] for r in rows], "has_header": True},
    ]


def _pptx_alt_text(shape) -> str:
    """画像の代替テキスト（PowerPointの「代替テキスト」欄）。画像そのものの中身は読み取れないため、
    作成者が書いた説明があればそれを使う（2026-10-01）"""
    descr = shape._element.xpath("./*/p:cNvPr/@descr")
    text = (descr[0] if descr else "").strip()
    if not text or any(m in text for m in _AUTO_ALT_TEXT_MARKERS):
        return ""
    # 作成ツールによっては代替テキスト欄に画像のファイル名（image.png等）が入るだけのため除外する
    if _IMAGE_FILE_NAME.fullmatch(text):
        return ""
    return text


def _pptx_shape_blocks(shape) -> list[dict]:
    """スライド上の図形1つ分のブロック。グループ図形は中の図形を再帰的に辿る"""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [b for child in shape.shapes for b in _pptx_shape_blocks(child)]
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return _pptx_chart_blocks(shape.chart)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        alt = _pptx_alt_text(shape)
        return [{"type": "paragraph", "text": f"画像: {alt}"}] if alt else []
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[{"text": cell.text, "span": 1} for cell in row.cells] for row in shape.table.rows]
        return [{"type": "table", "rows": rows}]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        blocks = []
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs) or para.text
            if not text.strip():
                continue
            # 本文プレースホルダーの段落は箇条書きとして書かれるのが普通なので、インデント付き
            # （level>0）だけでなく本文の段落もlist_itemとして扱うと実物の見た目に近くなる
            blocks.append({"type": "list_item" if para.level > 0 or shape.is_placeholder else "paragraph", "text": text})
        return blocks
    return []


def _pptx_to_blocks(data: bytes) -> dict:
    prs = Presentation(io.BytesIO(data))
    blocks: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks.append({"type": "heading", "level": 2, "text": f"スライド{i}" + (f": {title}" if title else "")})
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            blocks.extend(_pptx_shape_blocks(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append({"type": "paragraph", "text": f"ノート: {notes}"})
        if len(blocks) >= MAX_BLOCKS:
            return {"blocks": blocks[:MAX_BLOCKS], "truncated": True}
    return {"blocks": blocks, "truncated": False}


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour, v.minute, v.second) != (0, 0, 0) else v.strftime("%Y-%m-%d")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _xlsx_to_blocks(data: bytes, max_rows: int | None) -> dict:
    # data_only=Trueで数式は「最後にExcelで保存されたときの計算結果」を読む。Excel以外のツールで
    # 生成され計算結果が保存されていないファイルでは、数式のセルは空になる（openpyxlは数式を計算しない）
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[dict] = []
    truncated = False
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for raw in ws.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in raw[:XLSX_MAX_COLS]]
                if not any(cells):
                    continue
                if max_rows is not None and len(rows) >= max_rows:
                    truncated = True
                    break
                rows.append(cells)
            if not rows:
                continue
            # 右端の空列を落とし、行ごとの列数を揃える（書式だけ設定された列が延々と続くのを防ぐ）
            width = max(max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows)
            rows = [(r + [""] * width)[:width] for r in rows]
            blocks.append({"type": "heading", "level": 2, "text": f"シート: {ws.title}"})
            blocks.append({"type": "table", "rows": [[{"text": c, "span": 1} for c in r] for r in rows]})
    finally:
        wb.close()
    return {"blocks": blocks, "truncated": truncated}


def office_to_blocks(data: bytes, ext: str) -> dict:
    """プレビュー用のブロック列。extは小文字・ドット付き（OFFICE_EXTENSIONSのいずれか）"""
    if ext == ".docx":
        return _docx_to_blocks(data)
    if ext == ".pptx":
        return _pptx_to_blocks(data)
    if ext == ".xlsx":
        return _xlsx_to_blocks(data, XLSX_PREVIEW_MAX_ROWS)
    raise ValueError(f"unsupported: {ext}")


def _table_to_text(rows: list[list[dict]], labeled: bool) -> list[str]:
    """表を索引用のテキスト行にする。Excelのシートは1行目を見出しとみなし、各行を
    「見出し: 値 / 見出し: 値」の形にする（検索で一部の行だけがAIへ渡されても、どの列の値か
    分かるようにするため。行が文字数ベースのチャンク分割で途中で切れても同様）。
    スライド内の表は見出し行の有無が一定しないため、セルを「 | 」で区切るだけにする（グラフから
    取り出した表は1行目が必ず見出しのため、Excelと同じ形にする）"""
    texts = [[c["text"] for c in r] for r in rows]
    if not labeled or len(texts) < 2:
        return [" | ".join(t for t in r if t) for r in texts if any(r)]
    header = [h or get_column_letter(i + 1) for i, h in enumerate(texts[0])]
    lines = [" / ".join(h for h in texts[0] if h)]
    for r in texts[1:]:
        pairs = [f"{header[i]}: {v}" for i, v in enumerate(r) if v]
        if pairs:
            lines.append(" / ".join(pairs))
    return lines


def office_text_for_index(data: bytes, ext: str) -> str:
    """索引化（doc_indexer.py）用のテキスト。PowerPoint・Excelのみ（Wordは従来の抽出のまま）"""
    if ext == ".pptx":
        blocks = _pptx_to_blocks(data)["blocks"]
    elif ext == ".xlsx":
        blocks = _xlsx_to_blocks(data, None)["blocks"]
    else:
        raise ValueError(f"unsupported: {ext}")
    lines: list[str] = []
    for b in blocks:
        if b["type"] == "heading":
            lines.extend(["", f"【{b['text']}】"])
        elif b["type"] == "list_item":
            lines.append(f"・{b['text']}")
        elif b["type"] == "table":
            lines.extend(_table_to_text(b["rows"], labeled=ext == ".xlsx" or b.get("has_header", False)))
        else:
            lines.append(b["text"])
    return "\n".join(lines).strip()
